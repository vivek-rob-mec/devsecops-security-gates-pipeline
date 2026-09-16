"""Build editable PPTX, matching PDF, notes handbook and preview from one source.

Install requirements.txt first. No PowerPoint automation is required.
The PDF uses the same layout model; it is not a render by Microsoft PowerPoint.
"""
from pathlib import Path
import sys
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '.tools' / 'python'))
import json
import re
import shutil
import html
import pymupdf as fitz
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from course_content import SLIDES, SOURCES, REVIEWED

OUT = ROOT / 'presentation'
W, H = 960, 540
INK, MUTED, BG = '182F45', '486176', 'F4F7FA'
TEAL, GOLD, WHITE = '007E83', 'D09830', 'FFFFFF'
FONT_DIR = Path('C:/Windows/Fonts')
FONT_FILE = FONT_DIR / 'arial.ttf'
BOLD_FILE = FONT_DIR / 'arialbd.ttf'
if not FONT_FILE.exists():
    FONT_FILE = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
    BOLD_FILE = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf')
if not FONT_FILE.exists():
    raise SystemExit('Install Arial or DejaVu Sans to measure and export text.')
FONT_NAME = 'Arial' if FONT_FILE.name == 'arial.ttf' else 'DejaVu Sans'
prs = Presentation()
prs.slide_width, prs.slide_height = Pt(W), Pt(H)
pdf = fitz.open()
measurements = []

def rgb(value):
    return tuple(int(value[i:i+2], 16) / 255 for i in (0, 2, 4))

def wrap(text, width, size, bold=False):
    font = ImageFont.truetype(str(BOLD_FILE if bold else FONT_FILE), round(size*4))
    lines = []
    for paragraph in text.split('\n'):
        line = ''
        for word in paragraph.split():
            candidate = (line + ' ' + word).strip()
            if font.getlength(candidate)/4 > width:
                if not line:
                    raise ValueError(f'Unbreakable text exceeds width: {word}')
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines

def rect(slide, page, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    page.draw_rect(fitz.Rect(x,y,x+w,y+h), color=None, fill=rgb(color))

def text(slide, page, value, x, y, w, h, size=18, color=INK, bold=False, link=None):
    lines = wrap(value, w-2, size, bold)
    leading = size*1.24
    used = len(lines)*leading+3
    if used > h:
        raise ValueError(f'Slide {len(prs.slides)} overflow: {value[:65]} ({used:.1f}>{h})')
    if min(x,y)<0 or x+w>W+0.1 or y+h>H+0.1:
        raise ValueError('Shape outside slide')
    shape = slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False  # explicit measured line breaks make the layout stable
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        p = tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT_NAME
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(leading)
        if link and p.runs:
            p.runs[0].hyperlink.address = link
        page.insert_text((x,y+size+i*leading), line, fontsize=size,
                         fontname='Bold' if bold else 'Regular', color=rgb(color))
    if link:
        page.insert_link({'kind':fitz.LINK_URI,'from':fitz.Rect(x,y,x+w,y+h),'uri':link})
    measurements.append(dict(slide=len(prs.slides), text=value[:60], lines=len(lines), font=size, spare=round(h-used,1)))

def notes(item):
    result = item['notes']
    if item['sources']:
        result += '\n\nPrimary sources (reviewed '+REVIEWED+'):\n'
        result += '\n'.join(SOURCES[k][0]+': '+SOURCES[k][1] for k in item['sources'])
    return result

def make_slide(item, number, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    page = pdf.new_page(width=W,height=H)
    page.insert_font(fontname='Regular',fontfile=str(FONT_FILE))
    page.insert_font(fontname='Bold',fontfile=str(BOLD_FILE))
    rect(slide,page,0,0,W,H,BG)
    rect(slide,page,0,0,10,H,TEAL)
    text(slide,page,item['section'],34,23,870,20,11,TEAL,True)
    text(slide,page,item['title'],34,58,886,83,30,INK,True)
    text(slide,page,item['subtitle'],34,141,886,45,16,MUTED)
    count=len(item['cards'])
    gap=18 if count==3 else 23
    cardw=(886-gap*(count-1))/count
    for i,(heading,body) in enumerate(item['cards']):
        x=34+i*(cardw+gap)
        rect(slide,page,x,208,cardw,211,WHITE)
        rect(slide,page,x,208,cardw,4,TEAL if i%2==0 else GOLD)
        text(slide,page,heading,x+15,229,cardw-30,48,15,TEAL,True)
        text(slide,page,body,x+15,284,cardw-30,125,17 if count==3 else 16,INK)
        if item['layout']=='flow' and i<count-1:
            text(slide,page,'›',x+cardw+3,300,19,35,23,TEAL,True)
    rect(slide,page,34,442,886,45,INK)
    text(slide,page,item['takeaway'],46,453,860,28,14,WHITE,True)
    source_label = 'Sources in notes: '+', '.join(item['sources']) if item['sources'] else 'Project guidance • examples are illustrative'
    text(slide,page,source_label,34,505,804,20,9,MUTED)
    text(slide,page,f'{number:02d} / {total:02d}',855,503,68,22,11,MUTED,True)
    slide.notes_slide.notes_text_frame.text = notes(item)
    return page

def main():
    items=list(SLIDES)
    source_items=list(SOURCES.items())
    for start in range(0,len(source_items),6):
        chunk=source_items[start:start+6]
        cards=[]
        for offset in range(0,len(chunk),2):
            pair=chunk[offset:offset+2]
            cards.append(('PRIMARY REFERENCES', '\n\n'.join(v[0] for _,v in pair)))
        items.append(dict(section='09 SOURCES', title='Sources and maintenance',
            subtitle='Official documentation; full clickable links are in the slide notes and handbook.',
            cards=cards, notes='These references support the tool-specific descriptions. Review behavior against your pinned versions when maintaining the project. Product behavior and publication schedules can change.',
            sources=[k for k,_ in chunk],takeaway='Source review date: '+REVIEWED+'. Local configuration determines actual behavior.',layout='cards'))
    for i,item in enumerate(items,1):
        make_slide(item,i,len(items))
    dest=OUT/'DevSecOps_Industry_Standard_Pipeline.pptx'
    backup=OUT/'source-original'/'DevSecOps_Industry_Standard_Pipeline.before-revision.pptx'
    if dest.exists() and not backup.exists():
        shutil.copy2(dest,backup)
    prs.core_properties.title='DevSecOps: from first principles to release evidence'
    prs.core_properties.subject='Pipeline flow, security-tool internals and implementation boundaries'
    prs.core_properties.author='DevSecOps reference project'
    prs.save(dest)
    pdf.save(OUT/'DevSecOps_Industry_Standard_Pipeline.pdf',garbage=4,deflate=True)
    handbook=['# DevSecOps handbook\n',
        'Full offline companion to the presentation. Source review: '+REVIEWED+'.\n',
        '## How to use this handbook\n',
        'Read foundations first, then pipeline flow and detection. Intermediate readers can focus on intelligence and project practice. Professionals should inspect the release contract and evidence controls. Slide notes contain the same explanations. The separate implementation guides describe commands and current limits.\n',
        '## Contents\n']
    for i,item in enumerate(items,1):
        handbook.append(f'- [{i:02d}. {item["title"]}](#lesson-{i:02d})')
    for i,item in enumerate(items,1):
        handbook += [f'\n<a id="lesson-{i:02d}"></a>\n',f'## {i:02d}. {item["title"]}\n',f'*{item["section"]} — {item["subtitle"]}*\n']
        for heading,body in item['cards']:
            handbook.append(f'**{heading}.** '+body.replace('\n','; ')+'\n')
        handbook.extend([item['notes']+'\n', '**Key point:** '+item['takeaway']+'\n'])
        if item['sources']:
            handbook.append('Sources: '+', '.join(f'[{SOURCES[k][0]}]({SOURCES[k][1]})' for k in item['sources'])+'.\n')
    (ROOT/'docs'/'DEVSECOPS_HANDBOOK.md').write_text('\n'.join(handbook),encoding='utf-8')
    html_parts=['<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">',
        '<title>DevSecOps offline handbook</title><style>body{font:18px/1.7 system-ui,sans-serif;color:#182f45;background:#f4f7fa;max-width:1000px;margin:auto;padding:32px}a{color:#007e83}article{background:white;padding:30px;margin:32px 0;border-top:4px solid #007e83}h1,h2{line-height:1.2}h3{font-size:1em;margin-bottom:0}nav{columns:2}small{color:#486176}@media(max-width:650px){nav{columns:1}body{padding:16px}}@media print{body{background:white;font-size:11pt}nav{display:none}article{break-inside:avoid;padding:10px}}</style>',
        '<h1>DevSecOps offline handbook</h1><p>Full slide content, explanations and primary sources. Review date: '+REVIEWED+'. Use browser search to find a term; no network connection is needed to read the course.</p>',
        '<p><a href="../references/glossary.md">Expanded glossary</a> · <a href="implementation/quickstart.md">Implementation quickstart</a></p><nav><ol>']
    html_parts.extend(f'<li><a href="#lesson-{i:02d}">{html.escape(s["title"])}</a></li>' for i,s in enumerate(items,1))
    html_parts.append('</ol></nav>')
    for i,item in enumerate(items,1):
        html_parts.append(f'<article id="lesson-{i:02d}"><small>{html.escape(item["section"])}</small><h2>{i:02d}. {html.escape(item["title"])}</h2><p>{html.escape(item["subtitle"])}</p>')
        for heading,body in item['cards']:
            html_parts.append('<h3>'+html.escape(heading)+'</h3><p>'+html.escape(body).replace('\n','<br>')+'</p>')
        html_parts.append('<h3>Explanation</h3><p>'+html.escape(item['notes'])+'</p><p><strong>'+html.escape(item['takeaway'])+'</strong></p>')
        if item['sources']:
            html_parts.append('<p>Sources: '+', '.join('<a href="'+html.escape(SOURCES[k][1],quote=True)+'">'+html.escape(SOURCES[k][0])+'</a>' for k in item['sources'])+'</p>')
        html_parts.append('</article>')
    html_parts.append('</html>')
    (ROOT/'docs'/'DEVSECOPS_HANDBOOK.html').write_text('\n'.join(html_parts),encoding='utf-8')
    (OUT/'slide-index.md').write_text('# Slide index\n\n'+'\n'.join(f'{i}. **{s["section"]}** — {s["title"]}' for i,s in enumerate(items,1))+'\n',encoding='utf-8')
    preview=OUT/'preview'
    preview.mkdir(exist_ok=True)
    thumbs=[]
    for i,page in enumerate(pdf):
        pix=page.get_pixmap(matrix=fitz.Matrix(.5,.5))
        im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
        thumbs.append(im)
    columns=4
    sheet=Image.new('RGB',(480*columns,290*((len(thumbs)+columns-1)//columns)), '#DCE4EA')
    for i,im in enumerate(thumbs):
        sheet.paste(im,((i%columns)*480,(i//columns)*290))
    sheet.save(preview/'contact-sheet.png')
    for i in [0,7,21,26,40]:
        if i<len(pdf): pdf[i].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(preview/f'slide-{i+1:02d}.png')
    (preview/'layout-check.json').write_text(json.dumps({'slides':len(items),'text_boxes':len(measurements),'minimum_body_font':16,'checks':measurements},indent=2),encoding='utf-8')
    print(f'Built {len(items)} slides; {len(measurements)} text boxes fit measured bounds. PPTX, PDF and handbook saved.')

if __name__=='__main__': main()
